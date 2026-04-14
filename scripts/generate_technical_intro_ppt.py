"""Generate a technical SIAS introduction deck from the paper."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


PRIMARY = RGBColor(27, 54, 93)
ACCENT = RGBColor(18, 132, 198)
ACCENT_2 = RGBColor(74, 111, 165)
BG = RGBColor(245, 247, 250)
TEXT = RGBColor(34, 40, 49)
MUTED = RGBColor(102, 112, 133)
SUCCESS = RGBColor(34, 139, 34)
WARN = RGBColor(196, 112, 16)
WHITE = RGBColor(255, 255, 255)


def set_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_header(slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(11.8), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(27)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.95), Inches(11.4), Inches(0.35))
        stf = sub_box.text_frame
        p = stf.paragraphs[0]
        r = p.add_run()
        r.text = subtitle
        r.font.size = Pt(11)
        r.font.color.rgb = MUTED


def add_footer(slide, page: int) -> None:
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.6),
        Inches(6.92),
        Inches(11.7),
        Inches(0.03),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()

    page_box = slide.shapes.add_textbox(Inches(11.75), Inches(6.98), Inches(0.4), Inches(0.2))
    tf = page_box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = str(page)
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED


def add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    top_band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(1.25)
    )
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = PRIMARY
    top_band.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.72), Inches(1.55), Inches(11.8), Inches(1.2))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "SIAS 技术介绍"
    r.font.size = Pt(29)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "Streaming Importance-Aware Selection for Continual Agent Tool Use"
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT

    body = slide.shapes.add_textbox(Inches(0.92), Inches(3.0), Inches(11.0), Inches(2.2))
    tf = body.text_frame
    bullets = [
        "基于 paper 重构，面向技术人员讲清楚：问题定义、算法路径、实现映射、实验协议。",
        "核心目标：让 agent 在持续到来的 tool-use 任务流中适应新分布，同时尽量不遗忘旧能力。",
        "方法焦点：adaptive replay、streaming coreset、unified importance scoring。",
    ]
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT
        p.space_after = Pt(10)

    badge = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(5.72), Inches(4.6), Inches(0.58)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = ACCENT
    badge.line.fill.background()
    tf = badge.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Paper-Aligned Technical Deck"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = WHITE


def add_bullets_slide(
    prs: Presentation,
    page: int,
    title: str,
    subtitle: str,
    bullets: list[str],
    *,
    font_size: float = 20,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, title, subtitle)

    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.7), Inches(5.9))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(10)

    add_footer(slide, page)


def add_two_col_slide(
    prs: Presentation,
    page: int,
    title: str,
    subtitle: str,
    left_title: str,
    left_items: list[str],
    right_title: str,
    right_items: list[str],
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, title, subtitle)

    panels = [
        (Inches(0.72), RGBColor(236, 243, 251), left_title, left_items),
        (Inches(6.82), RGBColor(236, 248, 251), right_title, right_items),
    ]
    for x, fill, head, items in panels:
        panel = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, Inches(1.55), Inches(5.7), Inches(5.15)
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = fill
        panel.line.color.rgb = ACCENT_2
        panel.line.width = Pt(1.1)

        head_box = slide.shapes.add_textbox(x + Inches(0.22), Inches(1.75), Inches(5.0), Inches(0.35))
        tf = head_box.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = head
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = PRIMARY

        body_box = slide.shapes.add_textbox(x + Inches(0.22), Inches(2.15), Inches(5.08), Inches(4.2))
        tf = body_box.text_frame
        tf.word_wrap = True
        for idx, item in enumerate(items):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = item
            p.font.size = Pt(15.5)
            p.font.color.rgb = TEXT
            p.space_after = Pt(8)

    add_footer(slide, page)


def add_problem_slide(prs: Presentation, page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, "为什么要研究 Continual Agent Tool Use", "现有工具调用评测大多是 static/reset-after-each-task，和真实部署存在落差")

    left = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(1.8), Inches(5.35), Inches(3.65)
    )
    left.fill.solid()
    left.fill.fore_color.rgb = RGBColor(237, 243, 250)
    left.line.color.rgb = ACCENT_2
    tf = left.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Static Evaluation"
    r.font.size = Pt(21)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    for line in [
        "任务单次执行",
        "环境重置",
        "无持久记忆状态",
        "无法观测长期遗忘",
    ]:
        p = tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT

    right = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(1.8), Inches(5.35), Inches(3.65)
    )
    right.fill.solid()
    right.fill.fore_color.rgb = RGBColor(236, 248, 251)
    right.line.color.rgb = ACCENT
    tf = right.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Deployment Reality"
    r.font.size = Pt(21)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    for line in [
        "任务流持续到来",
        "工具和环境分布漂移",
        "新能力学习与旧能力保持冲突",
        "内存和重放预算受限",
    ]:
        p = tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT

    arrow = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(5.98), Inches(3.0), Inches(0.78), Inches(0.95)
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = WARN
    arrow.line.fill.background()

    note = slide.shapes.add_textbox(Inches(0.95), Inches(5.82), Inches(11.0), Inches(0.5))
    tf = note.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "SIAS 要解决的就是这个 gap：bounded memory 下的持续适应 + 抗遗忘。"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = TEXT

    add_footer(slide, page)


def add_pipeline_slide(prs: Presentation, page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, "SIAS 总体流程", "代码和论文一致的主路径：stream -> signals -> selector -> scorer -> replay/contract")

    steps = [
        ("Tool-use Stream", "instruction / tools / outcome"),
        ("Adaptive Learner", "drift + forgetting"),
        ("Streaming Selector", "short-term / long-term"),
        ("Importance Scorer", "loss + diversity + novelty"),
        ("Replay + Contract", "training batch + trace"),
    ]
    x = 0.55
    widths = [2.25, 2.35, 2.35, 2.15, 2.55]
    for idx, ((head, desc), width) in enumerate(zip(steps, widths)):
        box = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.25), Inches(width), Inches(1.35)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(233, 241, 249) if idx % 2 == 0 else RGBColor(234, 247, 250)
        box.line.color.rgb = ACCENT
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = head
        r.font.size = Pt(16.5)
        r.font.bold = True
        r.font.color.rgb = PRIMARY
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = desc
        r.font.size = Pt(11)
        r.font.color.rgb = MUTED
        if idx < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.CHEVRON,
                Inches(x + width + 0.1),
                Inches(2.68),
                Inches(0.42),
                Inches(0.48),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ACCENT
            arrow.line.fill.background()
        x += width + 0.52

    memory = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(3.9), Inches(4.45), Inches(5.45), Inches(1.15)
    )
    memory.fill.solid()
    memory.fill.fore_color.rgb = RGBColor(232, 245, 236)
    memory.line.color.rgb = SUCCESS
    tf = memory.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Bounded Replay Memory"
    r.font.size = Pt(19)
    r.font.bold = True
    r.font.color.rgb = SUCCESS
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "选中的高价值轨迹进入 buffer，并按 adaptive replay ratio 回流训练"
    r.font.size = Pt(12.5)
    r.font.color.rgb = TEXT

    add_footer(slide, page)


def add_formula_slide(
    prs: Presentation,
    page: int,
    title: str,
    subtitle: str,
    formula: str,
    explainer: list[str],
    impl_points: list[str],
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, title, subtitle)

    formula_box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(1.65), Inches(11.35), Inches(1.15)
    )
    formula_box.fill.solid()
    formula_box.fill.fore_color.rgb = RGBColor(233, 240, 249)
    formula_box.line.color.rgb = ACCENT
    tf = formula_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = formula
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = PRIMARY

    add_two_col_slide_content(
        slide,
        left_title="论文解释",
        left_items=explainer,
        right_title="代码落点",
        right_items=impl_points,
    )
    add_footer(slide, page)


def add_two_col_slide_content(slide, left_title: str, left_items: list[str], right_title: str, right_items: list[str]) -> None:
    panels = [
        (Inches(0.82), RGBColor(236, 243, 251), left_title, left_items),
        (Inches(6.82), RGBColor(236, 248, 251), right_title, right_items),
    ]
    for x, fill, head, items in panels:
        panel = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, Inches(3.15), Inches(5.6), Inches(3.3)
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = fill
        panel.line.color.rgb = ACCENT_2

        head_box = slide.shapes.add_textbox(x + Inches(0.2), Inches(3.32), Inches(4.9), Inches(0.3))
        tf = head_box.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = head
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = PRIMARY

        body = slide.shapes.add_textbox(x + Inches(0.2), Inches(3.7), Inches(4.95), Inches(2.45))
        tf = body.text_frame
        tf.word_wrap = True
        for idx, item in enumerate(items):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = item
            p.font.size = Pt(14.8)
            p.font.color.rgb = TEXT
            p.space_after = Pt(6)


def add_table_slide(
    prs: Presentation,
    page: int,
    title: str,
    subtitle: str,
    headers: list[str],
    rows: list[tuple[str, str, str]],
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, title, subtitle)

    col_x = [0.7, 3.05, 7.2]
    col_w = [2.15, 4.0, 5.1]
    y = 1.55
    for x, w, header in zip(col_x, col_w, headers):
        cell = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.52))
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        cell.line.fill.background()
        tf = cell.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = header
        r.font.size = Pt(14.5)
        r.font.bold = True
        r.font.color.rgb = WHITE

    y += 0.55
    fills = [RGBColor(240, 245, 251), RGBColor(235, 247, 250)]
    for idx, row in enumerate(rows):
        for x, w, text in zip(col_x, col_w, row):
            cell = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.88)
            )
            cell.fill.solid()
            cell.fill.fore_color.rgb = fills[idx % 2]
            cell.line.color.rgb = ACCENT_2
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = text
            r.font.size = Pt(12.5)
            r.font.color.rgb = TEXT
        y += 0.88

    add_footer(slide, page)


def add_impl_status_slide(prs: Presentation, page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, "当前实现状态", "这篇 paper 对应的是一个已经具备核心算法与实验骨架的代码库，而不是纯概念设计")

    done = [
        "Adaptive replay: drift_score、forgetting_score、adaptive_replay_ratio、策略切换",
        "Streaming coreset: streaming selector、short-term/long-term pools、benchmark 脚本",
        "Unified scoring: loss/diversity/novelty 统一打分和 importance_breakdown",
        "Benchmark contract: `sias-benchmark-v1`、配置对象、schema 校验",
        "Experiment framework: BFCL / AgentBench FC 的 phase 构造与 manifest 生成",
    ]
    pending = [
        "需要接入真实 BFCL 与 AgentBench FC 公共版本数据并跑完整 sweep",
        "需要补齐 EWC、MIR、DER++ 等 reviewer 预期 baseline",
        "需要生成最终论文表格：success、forgetting、BWT/FWT、runtime、memory",
    ]

    add_two_col_slide_content(slide, "已完成", done, "待补齐", pending)
    add_footer(slide, page)


def build_presentation(output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)
    add_problem_slide(prs, 2)
    add_bullets_slide(
        prs,
        3,
        "问题定义与目标",
        "论文把输入建模为非平稳的 tool-use experience stream",
        [
            "输入样本 x_t 包含：user instruction、tool schema(s)、intermediate execution context、outcome signal。",
            "数据流被分成多个 phase：可以按 BFCL 的 category/domain，或按 AgentBench FC 的 environment/task 构造。",
            "目标不是一次性离线最优，而是在 bounded memory 和 replay budget 下持续学习。",
            "评估重点包括：新阶段性能、旧阶段遗忘、forward/backward transfer、系统开销。",
        ],
    )
    add_pipeline_slide(prs, 4)
    add_formula_slide(
        prs,
        5,
        "模块一：Adaptive Replay",
        "用漂移和遗忘信号动态调 replay 强度，而不是固定 replay ratio",
        "r_t = clip(r_0 + alpha * max(d_t, f_t), r_min, r_max)",
        [
            "d_t 表示 distribution drift，f_t 表示 forgetting 强度，谁更强就由谁主导 replay 增强。",
            "如果 batch 更像“分布变了”，策略偏向 diversity；如果更像“已经开始忘”，策略偏向 loss_topk。",
            "核心思想是 replay 不应该固定，而要随着 stream 状态变化。",
        ],
        [
            "`OnlineContinualLearner` 维护 `drift_score`、`forgetting_score`、`adaptive_replay_ratio`。",
            "batch profile 由 topic/token/loss 统计构造，trace 持久化到 `metrics.jsonl`、`summary.csv`。",
            "训练 batch = new samples + replay samples，replay 由 selector 打分或随机抽样得到。",
        ],
    )
    add_formula_slide(
        prs,
        6,
        "模块二：Unified Importance Scoring",
        "buffer retention 和 replay selection 共用一套重要性定义，减少 heuristic 冲突",
        "I(x) = lambda_l * I_loss(x) + lambda_d * I_div(x) + lambda_n * I_nov(x)",
        [
            "loss 负责抓住困难样本，diversity 负责避免保留一堆相似轨迹，novelty 负责提升对新模式的覆盖。",
            "论文强调的不只是一个总分，而是 score breakdown 可解释。",
            "这种统一打分让 replay 和 coreset 不再各自用一套不一致的规则。",
        ],
        [
            "`ImportanceScorer` 默认权重为 loss 0.5 / diversity 0.3 / novelty 0.2。",
            "实现使用文本特征计算样本间相似度，并输出 `importance_breakdown`。",
            "contract 返回的不只是 selected sample，还包括 `importance_score` 与组件分解。",
        ],
    )
    add_two_col_slide(
        prs,
        7,
        "模块三：Streaming Coreset Maintenance",
        "避免每个 batch 都全量重算，从在线数据流里维护 bounded memory",
        "算法思路",
        [
            "维护 short-term pool 捕捉最近到来的样本；维护 long-term pool 保留历史高价值样本。",
            "新 batch 先进入 short-term pool，再与 long-term 合并形成 candidate pool。",
            "先按 importance 排出 long-term 候选，再用 hybrid 选择形成最终训练/保留集合。",
            "这样可以兼顾 recentness、coverage 和计算成本。",
        ],
        "代码实现",
        [
            "`CoresetSelector(strategy='streaming')` 暴露 `update_stream()`。",
            "支持 `streaming_window_size`、`replacement_temperature`、`long_term_pool_size`。",
            "内置 `loss_topk`、`diversity`、`hybrid`、`random`、`streaming` 五类策略。",
            "另有 `benchmark_streaming()` 直接测 selector 吞吐与池大小变化。",
        ],
    )
    add_table_slide(
        prs,
        8,
        "论文贡献与代码映射",
        "把 paper 中的贡献点直接映射到仓库模块，便于技术同学快速定位",
        ["贡献点", "主要模块", "技术含义"],
        [
            ("Adaptive replay", "`adaptive_signals.py` + `continual_learner.py`", "在线估计 drift/forgetting，并动态调 replay ratio 与 selection policy"),
            ("Streaming coreset", "`coreset_selector.py`", "在受限内存中做增量保留，不必每轮全量重算"),
            ("Unified scorer", "`importance_scorer.py`", "统一 replay 与 retention 的样本价值定义，并保留解释字段"),
            ("Stable contract", "`contract.py`", "把算法封装成 benchmark-facing schema，便于外部系统调用与复现"),
            ("Experiment scaffolding", "`experiment_framework.py`", "把真实 benchmark 数据转成 continual phases 与 manifest"),
        ],
    )
    add_two_col_slide(
        prs,
        9,
        "实验协议",
        "评测设计围绕 continual tool use，而不是普通静态 function calling",
        "主 benchmark: BFCL",
        [
            "覆盖 serial、parallel、multi-function、stateful multi-step function calling。",
            "可按 `test_category`、`category` 或 `domain` 构造 continual phases。",
            "主要看 tool-call success / pass rate，以及 phase 间遗忘与迁移。",
        ],
        "真实环境: AgentBench FC",
        [
            "把 OS、DB、KG、ALFWorld、WebShop 等环境转成 function-calling 风格任务。",
            "可按 `environment` 或 `task_name` 构造 phase。",
            "除了函数调用正确性，还能看 task completion 和跨环境保持能力。",
        ],
    )
    add_table_slide(
        prs,
        10,
        "Baseline 与指标",
        "技术汇报时建议把方法放到 continual learning 语境下，而不是只和普通 agent 比",
        ["维度", "选择", "目的"],
        [
            ("Baselines", "Sequential FT / Fixed Replay / Joint Oracle / EWC / MIR / DER++ / SIAS", "覆盖下界、上界、经典 replay、经典 regularization 和本文方法"),
            ("效果指标", "success/pass rate、task completion、forgetting、FWT、BWT", "同时衡量新任务适应和旧任务保持"),
            ("成本指标", "sample efficiency、runtime、memory overhead", "回答方法是否值得在在线系统里使用"),
            ("结果形式", "按 phase 报告曲线、最终汇总表、ablation", "让 reviewer 和工程同学都能看懂 trade-off"),
        ],
    )
    add_impl_status_slide(prs, 11)
    add_bullets_slide(
        prs,
        12,
        "给技术人员讲时建议强调什么",
        "这页可以直接当口播提纲",
        [
            "第一，SIAS 不是一个更大的 agent system，而是一个针对 continual tool use 的算法层。",
            "第二，它解决的是 replay 何时增强、哪些样本该留、为何留下来 这三个核心问题。",
            "第三，代码已经给出 benchmark-facing contract 和真实数据实验框架，所以后续接评测不是从零开始。",
            "第四，当前最大缺口在完整 benchmark sweep 和强 baseline 跑分，而不是方法定义本身。",
        ],
        font_size=18,
    )
    add_bullets_slide(
        prs,
        13,
        "结束页",
        "Takeaways",
        [
            "SIAS 把 continual learning 和 agent tool use 的交叉问题明确化，并给出可运行实现。",
            "技术核心是 adaptive replay + streaming coreset + unified importance scoring。",
            "论文结构、代码结构、实验协议三者已经基本对齐，适合继续推进真实 benchmark 结果产出。",
        ],
        font_size=19,
    )

    prs.save(output_path)


def main() -> None:
    repo = Path(__file__).resolve().parents[1]
    output_path = repo / "SIAS_Technical_Intro_From_Paper.pptx"
    build_presentation(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
