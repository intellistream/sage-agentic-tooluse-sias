"""Generate an updated SIAS presentation based on the current codebase."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


PRIMARY = RGBColor(27, 54, 93)
ACCENT = RGBColor(18, 132, 198)
ACCENT_2 = RGBColor(74, 111, 165)
BG = RGBColor(245, 247, 250)
TEXT = RGBColor(34, 40, 49)
MUTED = RGBColor(102, 112, 133)
SUCCESS = RGBColor(34, 139, 34)


def set_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_header(slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(11.6), Inches(0.7))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    p.alignment = PP_ALIGN.LEFT
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.95), Inches(11.2), Inches(0.4))
        stf = sub_box.text_frame
        p = stf.paragraphs[0]
        r = p.add_run()
        r.text = subtitle
        r.font.size = Pt(11)
        r.font.color.rgb = MUTED


def add_footer(slide, page: int) -> None:
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.6),
        Inches(6.95),
        Inches(11.7),
        Inches(0.03),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()
    box = slide.shapes.add_textbox(Inches(11.7), Inches(7.0), Inches(0.4), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = str(page)
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED


def add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(1.4)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = PRIMARY
    band.line.fill.background()

    title = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(11.8), Inches(1.5))
    tf = title.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "SIAS: Streaming Importance-Aware Selection for Continual Agent Tool Use"
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = PRIMARY

    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "根据当前代码库重构的论文版汇报 | 聚焦 agent tool 调用持续学习"
    r.font.size = Pt(18)
    r.font.color.rgb = ACCENT

    content = slide.shapes.add_textbox(Inches(0.9), Inches(3.1), Inches(11.2), Inches(2.0))
    tf = content.text_frame
    bullets = [
        "当前实现不再是早期“多专家 + 反射记忆”的概念原型，而是一个可运行的 L3 算法库。",
        "核心能力已经落地为：adaptive replay、streaming coreset、unified importance scorer、benchmark contract。",
        "论文主评测固定为 BFCL，真实环境评测固定为 AgentBench FC。",
    ]
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT

    badge = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.7), Inches(3.6), Inches(0.6)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = ACCENT
    badge.line.fill.background()
    tf = badge.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Code-Aligned Research Deck"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)


def add_bullets_slide(prs: Presentation, page: int, title: str, subtitle: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, title, subtitle)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.6), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(21 if i == 0 else 20)
        p.font.color.rgb = TEXT
        p.space_after = Pt(12)
    add_footer(slide, page)


def add_two_col_slide(
    prs: Presentation,
    page: int,
    title: str,
    subtitle: str,
    left_title: str,
    left_bullets: list[str],
    right_title: str,
    right_bullets: list[str],
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, title, subtitle)

    for x, head, bullets, fill in [
        (0.75, left_title, left_bullets, RGBColor(235, 242, 250)),
        (6.75, right_title, right_bullets, RGBColor(236, 247, 251)),
    ]:
        panel = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.55), Inches(5.7), Inches(5.2)
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = fill
        panel.line.color.rgb = ACCENT_2
        panel.line.width = Pt(1.2)
        head_box = slide.shapes.add_textbox(Inches(x + 0.2), Inches(1.75), Inches(5.1), Inches(0.4))
        tf = head_box.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = head
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = PRIMARY

        body = slide.shapes.add_textbox(Inches(x + 0.2), Inches(2.2), Inches(5.1), Inches(4.2))
        tf = body.text_frame
        tf.word_wrap = True
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = bullet
            p.level = 0
            p.font.size = Pt(15.5)
            p.font.color.rgb = TEXT
            p.space_after = Pt(8)
    add_footer(slide, page)


def add_pipeline_slide(prs: Presentation, page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, "方法总览", "当前代码中已经可运行的 SIAS 主路径")

    steps = [
        ("Input Stream", "agent trajectory / function-call sample"),
        ("Adaptive Learner", "drift_score / forgetting_score / adaptive_replay_ratio"),
        ("Streaming Selector", "short-term + long-term pools"),
        ("Unified Scorer", "loss + diversity + novelty"),
        ("Benchmark Contract", "sias-benchmark-v1 output"),
    ]
    x = 0.65
    widths = [2.0, 2.4, 2.2, 2.0, 2.2]
    for idx, ((title, subtitle), width) in enumerate(zip(steps, widths), 1):
        box = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.45), Inches(width), Inches(1.35)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(232, 240, 249) if idx % 2 else RGBColor(231, 246, 250)
        box.line.color.rgb = ACCENT
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = PRIMARY
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = subtitle
        r.font.size = Pt(11)
        r.font.color.rgb = MUTED
        if idx < len(steps):
            arrow = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(x + width + 0.1), Inches(2.82), Inches(0.45), Inches(0.55)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ACCENT
            arrow.line.fill.background()
        x += width + 0.55

    note = slide.shapes.add_textbox(Inches(0.85), Inches(4.5), Inches(11.3), Inches(1.8))
    tf = note.text_frame
    lines = [
        "当前 SIAS 代码库的核心贡献已经从“概念性 agent system”收敛为“面向 continual tool use 的算法层”。",
        "重点模块全部已经落地到 `src/sage_sias/`：`continual_learner.py`、`coreset_selector.py`、`importance_scorer.py`、`contract.py`。",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.font.color.rgb = TEXT
        p.space_after = Pt(8)
    add_footer(slide, page)


def add_table_like_slide(prs: Presentation, page: int, title: str, subtitle: str, rows: list[tuple[str, str, str]]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, title, subtitle)
    headers = ["模块", "当前实现", "论文意义"]
    col_x = [0.7, 3.2, 7.6]
    col_w = [2.3, 4.1, 4.9]
    y = 1.55
    for x, w, h in zip(col_x, col_w, headers):
        cell = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.55))
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        cell.line.fill.background()
        tf = cell.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = h
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    y += 0.58
    row_h = 0.9
    fills = [RGBColor(241, 245, 251), RGBColor(235, 247, 250)]
    for idx, row in enumerate(rows):
        for x, w, text in zip(col_x, col_w, row):
            cell = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(row_h)
            )
            cell.fill.solid()
            cell.fill.fore_color.rgb = fills[idx % 2]
            cell.line.color.rgb = ACCENT_2
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = text
            r.font.size = Pt(12.5)
            r.font.color.rgb = TEXT
        y += row_h
    add_footer(slide, page)


def add_roadmap_slide(prs: Presentation, page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, "投稿导向下一步", "当前代码已具备论文主线，接下来要补齐真实评测与正式 baseline")

    milestones = [
        ("Step 1", "接入真实 BFCL 数据", "按 `test_category/domain` 构造 continual phases，生成训练/评测 manifest。"),
        ("Step 2", "接入 AgentBench FC", "按 environment/task 构造长期任务流，验证真实环境中的遗忘与迁移。"),
        ("Step 3", "补传统 CL baseline", "增加 ER、EWC、MIR/DER++ runner，满足 reviewer 对 continual learning 对照的要求。"),
        ("Step 4", "产出论文图表", "success / forgetting / BWT / FWT / runtime / memory 六类核心图表统一导出。"),
    ]
    y = 1.55
    for idx, (step, head, desc) in enumerate(milestones):
        circ = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.85), Inches(y + 0.03), Inches(0.6), Inches(0.6)
        )
        circ.fill.solid()
        circ.fill.fore_color.rgb = ACCENT if idx < 2 else ACCENT_2
        circ.line.fill.background()
        tf = circ.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(idx + 1)
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

        box = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.65), Inches(y), Inches(10.6), Inches(0.72)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(237, 243, 250)
        box.line.color.rgb = ACCENT
        tf = box.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = f"{step} | {head}"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = PRIMARY
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = desc
        r.font.size = Pt(12)
        r.font.color.rgb = TEXT
        y += 1.15

    tag = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(6.2), Inches(4.8), Inches(0.55)
    )
    tag.fill.solid()
    tag.fill.fore_color.rgb = SUCCESS
    tag.line.fill.background()
    tf = tag.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Current status: algorithm code ready, benchmark integration next"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    add_footer(slide, page)


def build_presentation(output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)
    add_bullets_slide(
        prs,
        2,
        "为什么必须重做这版 PPT",
        "原始 PPT 与当前代码状态已经明显错位",
        [
            "早期版本强调 reflective memory、多专家协作、PPO/DPO 微调和自反式 agent system；这些内容并未形成当前仓库的主实现。",
            "当前代码真正完成的是一个可运行的 continual tool-use 算法栈：adaptive replay、streaming coreset、unified scoring、benchmark contract、real-dataset experiment framework。",
            "因此论文汇报必须从“系统概念展示”切换为“算法贡献 + 真实 benchmark 实验设计”。",
        ],
    )
    add_pipeline_slide(prs, 3)
    add_two_col_slide(
        prs,
        4,
        "核心模块 A：Adaptive Continual Learner",
        "对应 issue #5 / #11 已落地代码",
        "已实现能力",
        [
            "漂移/遗忘信号：`drift_score`、`forgetting_score`。",
            "动态重放：`adaptive_replay_ratio` 根据漂移和遗忘强度自动调整。",
            "策略切换：`loss_topk` / `diversity` / `hybrid` 规则式切换。",
            "可复现 trace：输出 `metrics.jsonl`、`summary.csv`、`config.snapshot.yaml`。",
        ],
        "论文价值",
        [
            "把 continual learning 从静态固定 replay 推进到 drift-aware replay。",
            "支持真实 tool-use 序列中的在线适应，而不是单轮离线选择。",
            "为 forgetting / sample efficiency / runtime 分析提供直接可观测日志。",
        ],
    )
    add_two_col_slide(
        prs,
        5,
        "核心模块 B：Streaming Coreset + Unified Scorer",
        "对应 issue #6 / #7 已落地代码",
        "已实现能力",
        [
            "Streaming selector 支持 `update_stream()`。",
            "短期 / 长期候选池与 reservoir-style 替换逻辑。",
            "统一评分器融合 `loss + diversity + novelty`。",
            "输出 `importance_score` 和 `importance_breakdown`，便于解释与调试。",
        ],
        "论文价值",
        [
            "解决 agent tool 调用场景下长期数据流的样本保留问题。",
            "比纯随机 replay 或仅基于 loss 的保留策略更可解释、更高效。",
            "为“为什么这些轨迹被保留/重放”提供结构化证据。",
        ],
    )
    add_table_like_slide(
        prs,
        6,
        "代码与论文贡献映射",
        "从 issue 完成状态映射到论文贡献点",
        [
            ("Adaptive Replay", "漂移检测、forgetting 检测、动态 replay_ratio、trace 输出", "核心方法贡献：减少遗忘、提升样本效率"),
            ("Streaming Coreset", "增量选择、短/长期池、基准脚本", "效率贡献：降低全量重算开销"),
            ("Unified Scorer", "loss/diversity/novelty 融合 + breakdown", "解释性贡献：统一 buffer 与 coreset 决策"),
            ("Contract / Integration", "v1 schema、adapter、integration helper", "工程贡献：可复现 benchmark 和上层集成"),
            ("Real Dataset Framework", "BFCL / AgentBench FC 适配与 manifest 生成", "实验贡献：接真实 benchmark 的桥梁"),
        ],
    )
    add_two_col_slide(
        prs,
        7,
        "主评测与真实环境评测",
        "根据你刚刚确定的投稿配置",
        "主评测：BFCL",
        [
            "定位：function calling / agentic evaluation 主 benchmark。",
            "优势：覆盖 serial、parallel、stateful multi-step setting。",
            "在我们的框架里：按 `test_category/domain` 构造成 continual phases。",
            "主指标：tool-call success、pass rate、forgetting、forward/backward transfer。",
        ],
        "真实环境：AgentBench FC",
        [
            "定位：真实 function-calling environment benchmark。",
            "环境：OS、DB、KG、ALFWorld、WebShop 等多任务场景。",
            "在我们的框架里：按 environment/task 构造成长期 skill stream。",
            "主指标：task completion、环境成功率、旧技能保持、新环境迁移能力。",
        ],
    )
    add_table_like_slide(
        prs,
        8,
        "Baseline 设计",
        "围绕 continual agent tool use 的两轴对照",
        [
            ("Tool-use baseline", "BFCL / AgentBench FC 原生 function-calling 评测", "证明方法在真实工具调用 benchmark 上有效"),
            ("Sequential FT", "不使用 replay / memory / regularization", "最基本下界"),
            ("Joint Oracle", "所有阶段混合训练", "经验上界"),
            ("Fixed Replay / ER", "固定 replay ratio 或标准经验回放", "直接对照 adaptive replay"),
            ("EWC / MIR / DER++", "传统 continual learning 强 baseline", "满足顶会 reviewer 对 CL 对照的预期"),
            ("SIAS", "adaptive replay + streaming coreset + unified scorer", "论文主方法"),
        ],
    )
    add_two_col_slide(
        prs,
        9,
        "真实数据实验框架",
        "当前代码已经可以把 BFCL / AgentBench FC 转成 continual benchmark",
        "已落地代码",
        [
            "`src/sage_sias/experiment_framework.py`：数据适配、phase 构造、manifest 生成。",
            "`scripts/run_real_dataset_experiment.py`：统一命令行入口。",
            "支持方法：`sequential_ft`、`fixed_replay`、`sias`、`joint_oracle`。",
            "自动输出：每 phase 的 train/eval JSONL、`manifest_index.csv`、`README.generated.md`。",
        ],
        "下一步只差真实数据接入",
        [
            "把真实 BFCL JSON/JSONL 样例字段对齐到适配器。",
            "把真实 AgentBench FC trajectory/log 对齐到 environment adapter。",
            "补 ER / EWC / MIR / DER++ runner。",
            "补论文表格与图表导出脚本。",
        ],
    )
    add_table_like_slide(
        prs,
        10,
        "当前仓库已完成的实验能力",
        "不是空白设计，而是已经有可执行骨架",
        [
            ("Ablation", "`run_tooluse_ablation.py` 生成 markdown/csv 报告", "支持无 replay / 固定 replay / 动态 replay / batch vs streaming coreset"),
            ("Streaming Benchmark", "`benchmark_streaming_selector.py`", "支持 batch recompute vs streaming selector 开销对比"),
            ("Contract Test", "`test_issue10_benchmark_contract.py`", "保障 schema 稳定性"),
            ("Integration Test", "`test_issue9_integration_contract.py`", "验证上层调用链"),
            ("Real Dataset Test", "`test_experiment_framework_real_datasets.py`", "验证 BFCL / AgentBench FC manifest 生成"),
        ],
    )
    add_bullets_slide(
        prs,
        11,
        "论文主张应该如何表述",
        "把叙事从“系统愿景”收敛为“算法贡献 + benchmark 证据”",
        [
            "问题：agent tool 调用在长序列任务中会遗忘旧工具使用模式，且面对新工具/新环境时容易产生适应与保持的冲突。",
            "方法：SIAS 通过 adaptive replay、streaming coreset、unified importance scoring，在持续工具调用场景中保留关键经验并减少无效重放。",
            "证据：在 BFCL 与 AgentBench FC 上，相比 Sequential FT、ER、EWC、MIR/DER++，实现更低 forgetting、更高 task success、更好的 sample efficiency。",
        ],
    )
    add_bullets_slide(
        prs,
        12,
        "当前结论",
        "这版代码已经足够支撑论文方法章节与实验设计章节",
        [
            "方法章节：已经有明确模块边界、主算法路径、输入输出契约和解释字段。",
            "实验章节：已经锁定 BFCL + AgentBench FC，并完成真实数据实验框架骨架。",
            "剩余最大工作量不在算法，而在真实数据清洗、baseline runner 补齐和正式图表产出。",
        ],
    )
    add_roadmap_slide(prs, 13)
    add_bullets_slide(
        prs,
        14,
        "结束页",
        "谢谢",
        [
            "SIAS 当前已经从概念原型转为可运行的 continual agent tool-use 算法库。",
            "接下来以 BFCL + AgentBench FC 为核心，补齐真实 benchmark 与传统 continual baseline，即可进入顶会投稿打磨阶段。",
        ],
    )

    prs.save(output_path)


def main() -> None:
    repo = Path(__file__).resolve().parents[1]
    output_path = repo / "SIAS_Updated_Code_Aligned_Presentation.pptx"
    build_presentation(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
